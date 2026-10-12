"""Visible codex-chrome UI acceptance using synthetic files and the real Rust server.

Buttons/fields use real browser interactions. CDP supplies a synthetic File to
the browser chooser and observes real exported Blobs; backend results are never
mocked. PDF/DOCX/image outputs are independently checked with pypdf/python-docx/PIL.
"""
import base64
import hashlib
import io
import json
import os
from pathlib import Path
import shlex
import time
import traceback
import urllib.request
import zipfile

from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "test-output/suite-media-ui"
OUT.mkdir(parents=True, exist_ok=True)
URL = os.environ.get("WORD_TEST_URL", "http://127.0.0.1:18768/")
CONFIG = dict(os.environ)
for line in (Path.home() / ".config/codex-chrome-mcp/env").read_text().splitlines():
    line = line.strip().removeprefix("export ")
    if line and not line.startswith("#") and "=" in line:
        key, value = line.split("=", 1)
        tokens = shlex.split(value)
        if tokens:
            CONFIG.setdefault(key, tokens[0])
BASE = CONFIG["CODEX_CHROME_API_URL_LAN"]
OPENER = urllib.request.build_opener(urllib.request.ProxyHandler({}))
TAB = None
CHECKS = []
EXPORTS = []
OWN_MATERIAL = None
ERROR = None
PREFIX = "suite-media-" + str(int(time.time()))


def call(name, **arguments):
    if TAB and name != "chrome_new_tab":
        arguments.setdefault("tab_id", TAB)
    request = urllib.request.Request(BASE + "/api/tools/call",
        data=json.dumps({"name": name, "arguments": arguments}).encode(),
        headers={"Authorization": "Bearer " + CONFIG["CODEX_CHROME_API_TOKEN"], "Content-Type": "application/json"})
    result = json.load(OPENER.open(request, timeout=45))
    if not result.get("ok") or result.get("result", {}).get("isError"):
        raise RuntimeError(str(result))
    raw = "".join(block.get("text", "") for block in result["result"]["content"] if block["type"] == "text")
    return json.loads(raw)


def evaluate(expression):
    result = call("chrome_cdp", method="Runtime.evaluate",
                  params={"expression": expression, "returnByValue": True, "awaitPromise": True})
    if result.get("exceptionDetails"):
        raise AssertionError(result["exceptionDetails"])
    return result.get("result", {}).get("value")


def click(selector):
    call("chrome_click", selector=selector)


def type_text(selector, value, clear=True):
    if clear and value == "":
        click(selector)
        call("chrome_keypress", key="Control+A")
        call("chrome_keypress", key="Backspace")
    else:
        call("chrome_type", selector=selector, text=value, clear=clear)


def select_option(selector, value):
    index = evaluate("[...document.querySelector(" + json.dumps(selector) + ").options].findIndex(o=>o.value===" + json.dumps(value) + ")")
    assert index >= 0
    click(selector)
    call("chrome_keypress", key="Home")
    for _ in range(index):
        call("chrome_keypress", key="ArrowDown")
    call("chrome_keypress", key="Enter")
    assert evaluate("document.querySelector(" + json.dumps(selector) + ").value") == value


def wait_for(expression, seconds=30):
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        value = evaluate(expression)
        if value:
            return value
        time.sleep(.4)
    raise AssertionError("Timed out waiting for " + expression)


def idle(seconds=30):
    wait_for("!document.querySelector('#suite-overlay').classList.contains('busy')", seconds)
    status = evaluate("({text:document.querySelector('#suite-status').textContent,error:document.querySelector('#suite-status').classList.contains('error')})")
    assert not status["error"], status
    return status["text"]


def check(condition, description):
    assert condition, description
    CHECKS.append(description)
    print("PASS " + description, flush=True)


def screenshot(name):
    result = call("chrome_cdp", method="Page.captureScreenshot", params={"format": "png"})
    (OUT / name).write_bytes(base64.b64decode(result["data"]))


def model():
    return evaluate("window.rustWord.document()")


def choose_synthetic_file(button, input_selector, path, mime):
    # The real button opens the picker promise. A synthetic File enters its
    # actual onchange/FileReader path, then the real Rust decoder is invoked.
    click(button)
    encoded = base64.b64encode(path.read_bytes()).decode()
    expression = "(() => {const input=document.querySelector(" + json.dumps(input_selector) + ");const data=Uint8Array.from(atob(" + json.dumps(encoded) + "),c=>c.charCodeAt(0));const files=new DataTransfer();files.items.add(new File([data]," + json.dumps(path.name) + ",{type:" + json.dumps(mime) + "}));input.files=files.files;input.dispatchEvent(new Event('change',{bubbles:true}));return input.files.length;})()"
    assert evaluate(expression) == 1
    idle(90)


def install_download_observer():
    # Observe the genuine Blob, retain original URL and click behavior, and let
    # Chrome perform its real download. No feature response is replaced.
    evaluate("""(() => {
      const urls=new Map(), create=URL.createObjectURL.bind(URL), click=HTMLAnchorElement.prototype.click;
      window.__suiteMediaExports=[];
      URL.createObjectURL=blob=>{const url=create(blob);if(blob instanceof Blob)urls.set(url,blob);return url;};
      HTMLAnchorElement.prototype.click=function(...args){
        if(this.download&&urls.has(this.href)){const name=this.download,blob=urls.get(this.href);
          blob.arrayBuffer().then(buffer=>{let raw='';const bytes=new Uint8Array(buffer);for(let i=0;i<bytes.length;i+=16384)raw+=String.fromCharCode(...bytes.subarray(i,i+16384));window.__suiteMediaExports.push({name,base64:btoa(raw),bytes:bytes.length});});
        }
        return click.apply(this,args);
      };return true;
    })()""")


def exported_after(selector, label):
    before = evaluate("window.__suiteMediaExports.length")
    click(selector)
    idle(90)
    value = wait_for("window.__suiteMediaExports.length>" + str(before) + " && window.__suiteMediaExports.at(-1)")
    data = base64.b64decode(value["base64"])
    destination = OUT / (label + Path(value["name"]).suffix)
    destination.write_bytes(data)
    EXPORTS.append({"name": value["name"], "bytes": len(data), "sha256": hashlib.sha256(data).hexdigest(), "path": str(destination)})
    return data


def save_open(path):
    before = model()
    saved = evaluate("window.rustWord.invoke('savePath',{path:" + json.dumps(path) + ",document:window.rustWord.document()})")
    opened = evaluate("window.rustWord.invoke('openPath',{path:" + json.dumps(path) + "})")
    check(opened["document"] == before, "真实 DOCX 保存重开保留完整模型：" + path)
    check(Path(saved["path"]).is_file() and zipfile.is_zipfile(saved["path"]), "Rust 实际写入磁盘 Word 文件：" + path)
    evaluate("window.rustWord.loadDocument(" + json.dumps(opened, ensure_ascii=False) + ")")
    return saved, opened


# Fixtures contain only synthetic text/images.
png_path = OUT / (PREFIX + ".png")
image = Image.new("RGB", (120, 80), "white")
draw = ImageDraw.Draw(image)
draw.rectangle((0, 0, 59, 39), fill=(211, 48, 59))
draw.rectangle((60, 0, 119, 39), fill=(30, 153, 87))
draw.rectangle((0, 40, 59, 79), fill=(40, 94, 190))
image.save(png_path)
pdf_path = OUT / (PREFIX + ".pdf")
writer = canvas.Canvas(str(pdf_path), pagesize=(420, 595))
for title in ["ALPHA PAGE", "BRAVO PAGE"]:
    writer.setFont("Helvetica-Bold", 30); writer.drawString(40, 510, title)
    writer.setFont("Helvetica", 16); writer.drawString(40, 460, "Synthetic browser acceptance document.")
    writer.showPage()
writer.save()
scan_path = OUT / (PREFIX + "-scanned.pdf")
scan = Image.new("RGB", (1400, 900), "white")
scan_draw = ImageDraw.Draw(scan)
scan_draw.text((100, 160), "扫描文档转可编辑文字", fill=(24, 40, 64), font=ImageFont.truetype("/System/Library/Fonts/STHeiti Medium.ttc", 72))
scan_draw.text((100, 360), "Invoice 20261012", fill=(24, 40, 64), font=ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 72))
writer = canvas.Canvas(str(scan_path), pagesize=(700, 450))
writer.drawImage(ImageReader(scan), 0, 0, width=700, height=450)
writer.showPage(); writer.save()

try:
    call("chrome_status")
    TAB = call("chrome_new_tab", url=URL)["tab"]["id"]
    (OUT / "tab-id.txt").write_text(TAB)
    wait_for("!!window.rustWord?.suite")
    if evaluate("!document.querySelector('#modal').classList.contains('hidden')"):
        check(evaluate("document.querySelector('#modal-title').textContent.includes('恢复')"), "隔离启动提示只包含测试草稿恢复")
        click("#modal-actions button:first-child")
    call("chrome_cdp", method="Page.setInterceptFileChooserDialog", params={"enabled": True})
    install_download_observer()
    initial = {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "媒体验收正文：图片与素材必须写入可编辑 Word。"}]}]}
    evaluate("window.rustWord.loadDocument(" + json.dumps({"document": initial, "name": "媒体验收 Word"}, ensure_ascii=False) + ")")
    click('[data-tab="tools"]')
    click('[data-action="imageTools"]')
    wait_for("document.querySelector('#suite-title').textContent==='图片编辑'")
    choose_synthetic_file("#image-open", "#image-input", png_path, "image/png")
    check(evaluate("window.rustWord.suite.getImage().width===120 && window.rustWord.suite.getImage().height===80 && document.querySelector('#image-preview').naturalWidth===120"), "图片打开经真实 FileReader/Rust 解码后显示原像素")
    click("#image-rotate"); idle()
    check(evaluate("window.rustWord.suite.getImage().width===80 && window.rustWord.suite.getImage().height===120"), "旋转按钮实际改变图片像素尺寸")
    for selector, value in [("#crop-x", "10"), ("#crop-y", "20"), ("#crop-width", "50"), ("#crop-height", "70")]:
        type_text(selector, value)
    click("#image-crop"); idle()
    check(evaluate("window.rustWord.suite.getImage().width===50 && window.rustWord.suite.getImage().height===70"), "裁剪按钮按当前像素坐标生成新图片")
    click("#image-undo"); idle()
    check(evaluate("window.rustWord.suite.getImage().width===80 && window.rustWord.suite.getImage().height===120"), "撤销图片裁剪恢复旋转后的真实图片")
    click("#image-undo"); idle()
    check(evaluate("window.rustWord.suite.getImage().width===120 && window.rustWord.suite.getImage().height===80"), "再次撤销恢复原图")
    click("#image-rotate"); idle()
    click("#image-crop"); idle()
    screenshot("image-edit.png")
    select_option("#image-format", "jpeg")
    exported = exported_after("#image-save", "image-export")
    decoded = Image.open(io.BytesIO(exported))
    check(decoded.format == "JPEG" and decoded.size == (50, 70), "图片导出按钮返回可独立读取的实际 JPEG")
    before_materials = evaluate("window.rustWord.invoke('materialList')")["materials"]
    click("#image-library"); idle()
    after_materials = evaluate("window.rustWord.invoke('materialList')")["materials"]
    own = next(m for m in after_materials if m["name"] == png_path.name and m["id"] not in {v["id"] for v in before_materials})
    OWN_MATERIAL = own["id"]
    check(own["width"] == 50 and own["height"] == 70, "编辑结果通过 UI 真正保存为本地素材")
    click("#image-insert"); idle()
    check(evaluate("document.querySelectorAll('.ProseMirror img').length") == 1, "图片编辑结果经插入按钮进入 Word 正文")
    click("#suite-close")
    click('[data-action="materials"]')
    idle()
    type_text("#material-search", png_path.name)
    wait_for("document.querySelectorAll('.material-card').length===1")
    click(".material-card button:nth-of-type(1)"); idle()
    check(evaluate("document.querySelector('.material-card img').naturalWidth===50"), "素材预览读取持久化图片文件")
    click(".material-card button:nth-of-type(2)"); idle()
    check(evaluate("document.querySelectorAll('.ProseMirror img').length") == 2, "素材插入按钮添加第二张实际嵌入图片")
    click("#suite-close")
    saved, opened = save_open("suite-media-before-import.docx")
    independent = Document(saved["path"])
    check(len(independent.inline_shapes) == 2 and "媒体验收正文" in "\n".join(p.text for p in independent.paragraphs), "python-docx 独立读取图文 Word，正文和两张图片完整")
    type_text(".ProseMirror > p:first-of-type", " 未保存保护原文", clear=False)
    old_model = model()
    check("未保存保护原文" in json.dumps(old_model, ensure_ascii=False), "真实键盘输入产生未保存 Word 修改")

    click('[data-tab="tools"]')
    click('[data-action="pdfTools"]')
    choose_synthetic_file("#pdf-open", "#pdf-input", pdf_path, "application/pdf")
    check(evaluate("document.querySelectorAll('.pdf-page').length===2 && [...document.querySelectorAll('.pdf-page img')].every(i=>i.naturalWidth>0)"), "PDF 实际打开两页，PDFKit 生成真实可见缩略图")
    click(".pdf-page:first-child button:nth-of-type(2)")
    check(evaluate("document.querySelector('.pdf-page:first-child span').textContent.includes('第 2 页')"), "PDF 下移按钮交换页序")
    click(".pdf-page:first-child button:nth-of-type(3)")
    check(evaluate("document.querySelector('.pdf-page:first-child span').textContent.includes('90°')"), "PDF 旋转按钮更新当前页面方向")
    type_text("#pdf-note-page", "2")
    type_text("#pdf-note-y", "35")
    type_text("#pdf-note", "浏览器复核批注")
    pdf_bytes = exported_after("#pdf-export", "pdf-edited")
    pdf = PdfReader(io.BytesIO(pdf_bytes))
    check(len(pdf.pages) == 2 and pdf.pages[0].get("/Rotate") == 90 and "BRAVO PAGE" in pdf.pages[0].extract_text(), "导出 PDF 保留真实重排与 90° 旋转")
    check(any(note.get_object().get("/Contents") == "浏览器复核批注" for note in pdf.pages[1].get("/Annots", [])), "导出文件含真正的标准文字批注对象")
    screenshot("pdf-edit.png")
    type_text("#pdf-range", "2")
    split = PdfReader(io.BytesIO(exported_after("#pdf-split", "pdf-split")))
    check(len(split.pages) == 1 and "BRAVO PAGE" in split.pages[0].extract_text(), "拆分按钮导出指定原始页为一页 PDF")
    image_zip = exported_after("#pdf-images", "pdf-images")
    with zipfile.ZipFile(io.BytesIO(image_zip)) as archive:
        pictures = [name for name in archive.namelist() if name.endswith(".png")]
        converted = Image.open(io.BytesIO(archive.read(pictures[0])))
        check(len(pictures) == 1 and converted.width > 600, "PDF 转图片按钮生成真实 PNG ZIP")
    click("#pdf-text"); idle(90)
    extracted = evaluate("document.querySelector('#pdf-text-result').value")
    check("ALPHA PAGE" in extracted and "BRAVO PAGE" in extracted and extracted.index("BRAVO PAGE") < extracted.index("ALPHA PAGE"), "提取文字按钮读取 PDF 原文本层并遵循页面顺序")
    type_text("#pdf-text-result", " 浏览器已校对", clear=False)
    click("#pdf-word")
    wait_for("!document.querySelector('#modal').classList.contains('hidden') && document.querySelector('#modal-title').textContent==='保存你的修改？'")
    check(model() == old_model, "PDF 导入前未保存 Word 保持原文等待用户决定")
    click("#modal-actions button:first-child")
    idle()
    check(model() == old_model and "已取消导入" in evaluate("document.querySelector('#suite-status').textContent"), "取消 PDF 导入保留未保存 Word 完整图文")
    click("#pdf-word")
    wait_for("!document.querySelector('#modal').classList.contains('hidden')")
    click("#modal-actions .primary")
    idle()
    check("BRAVO PAGE" in json.dumps(model()) and "浏览器已校对" in json.dumps(model(), ensure_ascii=False), "选择保存后 PDF 文字成为可编辑 Word 正文")
    old_saved = evaluate("window.rustWord.invoke('openPath',{path:'suite-media-before-import.docx'})")
    check("未保存保护原文" in json.dumps(old_saved["document"], ensure_ascii=False), "PDF 导入 guard 保存分支把旧图文修改真正写入原 DOCX")
    click("#suite-close")
    type_text(".ProseMirror > p:first-of-type", " 转换后真实编辑", clear=False)
    new_saved, _ = save_open("suite-media-from-pdf.docx")
    converted = Document(new_saved["path"])
    converted_text = "\n".join(p.text for p in converted.paragraphs)
    check("BRAVO PAGE" in converted_text and "ALPHA PAGE" in converted_text and "转换后真实编辑" in converted_text, "PDF 文字转换后经真实输入再次保存为标准可编辑 DOCX")

    # Raster-only PDF: no text layer succeeds only when the OCR UI is enabled.
    stable_word = model()
    click('[data-tab="tools"]'); click('[data-action="pdfTools"]')
    choose_synthetic_file("#pdf-open", "#pdf-input", scan_path, "application/pdf")
    click("#pdf-text")
    wait_for("!document.querySelector('#suite-overlay').classList.contains('busy')")
    check(evaluate("document.querySelector('#suite-status').classList.contains('error') && document.querySelector('#suite-status').textContent.includes('OCR')"), "扫描件无文本层时明确提示开启 OCR")
    click("#pdf-ocr"); click("#pdf-text"); idle(120)
    scanned_text = evaluate("document.querySelector('#pdf-text-result').value")
    check("扫描文档" in scanned_text and "20261012" in scanned_text, "OCR 勾选后实际识别中英文扫描件")
    check(model() == stable_word, "扫描件 OCR 预览不会覆盖当前 Word")
    screenshot("pdf-ocr.png")
    click("#suite-close")

    # Delete only this test's synthetic material via the actual confirmation UI.
    click('[data-action="materials"]'); idle()
    type_text("#material-search", png_path.name)
    wait_for("document.querySelectorAll('.material-card').length===1")
    click(".material-card button:nth-of-type(3)")
    wait_for("document.querySelector('#modal-title').textContent==='删除素材？'")
    click("#modal-actions button:first-child")
    idle()
    check(any(m["id"] == OWN_MATERIAL for m in evaluate("window.rustWord.invoke('materialList')")["materials"]), "取消素材删除保留实际文件")
    click(".material-card button:nth-of-type(3)")
    wait_for("document.querySelector('#modal-title').textContent==='删除素材？'")
    click("#modal-actions .primary"); idle()
    check(not any(m["id"] == OWN_MATERIAL for m in evaluate("window.rustWord.invoke('materialList')")["materials"]), "确认后只删除本次测试素材")
    OWN_MATERIAL = None
    click("#suite-close")
    screenshot("word-from-pdf.png")
except Exception:
    ERROR = traceback.format_exc()
    print(ERROR, flush=True)
    if TAB:
        try:
            screenshot("failure.png")
        except Exception:
            pass
finally:
    if TAB:
        if OWN_MATERIAL:
            try:
                evaluate("window.rustWord.invoke('materialDelete',{materialId:" + json.dumps(OWN_MATERIAL) + "})")
            except Exception:
                pass
        try:
            call("chrome_close_tab", tab_id=TAB)
        except Exception:
            pass
    (OUT / "results.json").write_text(json.dumps({"url": URL, "tab": TAB, "checks": CHECKS,
        "exports": EXPORTS, "error": ERROR}, ensure_ascii=False, indent=2))
if ERROR:
    raise SystemExit(1)
print("浏览器媒体验收全部通过：" + str(len(CHECKS)) + " 项。", flush=True)
