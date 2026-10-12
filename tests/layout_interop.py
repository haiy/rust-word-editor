"""Validate real OOXML layout with independently generated python-docx files.

Run with the bundled Python (python-docx is required) after cargo build.
Every roundtrip import removes word/rustword.json first.
"""
import json
from pathlib import Path
import subprocess
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Mm, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "test-output" / "layout"
OUT.mkdir(parents=True, exist_ok=True)
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def rpc(requests, allow_errors=False):
    encoded = "\n".join(json.dumps(dict(id=i, **r), ensure_ascii=False)
                        for i, r in enumerate(requests)) + "\n"
    process = subprocess.run([str(ROOT / "target/debug/rust-word"), "--rpc", str(OUT)],
                             input=encoded, text=True, capture_output=True, check=True)
    responses = [json.loads(line) for line in process.stdout.splitlines()]
    assert len(responses) == len(requests), process.stdout
    if allow_errors:
        return responses
    assert all("error" not in r for r in responses), responses
    return [r["data"] for r in responses]


def strip_snapshot(source, destination):
    with ZipFile(OUT / source) as zin, ZipFile(OUT / destination, "w", ZIP_DEFLATED) as zout:
        for entry in zin.infolist():
            if entry.filename != "word/rustword.json":
                zout.writestr(entry.filename, zin.read(entry.filename))


def near(actual, expected, tolerance=0.03):
    assert abs(actual - expected) < tolerance, (actual, expected)


def configure_a4(section, landscape=False):
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (Mm(297), Mm(210)) if landscape else (Mm(210), Mm(297))
    section.top_margin, section.bottom_margin = Mm(24), Mm(31)
    section.left_margin, section.right_margin = Mm(17), Mm(27)


external = Document()
configure_a4(external.sections[0], landscape=True)
normal = external.styles["Normal"].paragraph_format
normal.keep_with_next = normal.keep_together = normal.page_break_before = True
style = external.styles.add_style("Layout inherited", WD_STYLE_TYPE.PARAGRAPH)
style.base_style = external.styles["Normal"]
style.paragraph_format.space_before = Pt(12.5)
style.paragraph_format.space_after = Pt(6.5)
style.paragraph_format.left_indent = Pt(36)
style.paragraph_format.first_line_indent = Pt(-18)
p = external.add_paragraph("继承段落样式，但显式关闭分页开关。", style=style)
p.paragraph_format.keep_with_next = False
p.paragraph_format.keep_together = False
p.paragraph_format.page_break_before = False
p.paragraph_format.line_spacing = 1.75
p = external.add_paragraph("首行缩进并与下段保持在同一页。")
p.paragraph_format.first_line_indent = Pt(24)
p.paragraph_format.space_before = Pt(3)
p.paragraph_format.space_after = Pt(14)
external.save(OUT / "foreign-layout.docx")

loaded = rpc([{"command": "openPath", "path": "foreign-layout.docx"}])[0]
assert not loaded["protected"], loaded["warnings"]
document = loaded["document"]
assert document["attrs"]["orientation"] == "landscape"
for key, expected in [("marginTop", 24), ("marginBottom", 31), ("marginLeft", 17), ("marginRight", 27)]:
    near(document["attrs"][key], expected)
a, b = [p["attrs"] for p in document["content"]]
assert a["spaceBefore"] == 12.5 and a["spaceAfter"] == 6.5
assert a["firstLineIndent"] == -18 and a["indent"] == 2 and float(a["lineHeight"]) == 1.75
assert all(a[key] is False for key in ["keepNext", "keepLines", "pageBreakBefore"])
assert all(b[key] is True for key in ["keepNext", "keepLines", "pageBreakBefore"])
assert b["firstLineIndent"] == 24
rpc([{"command": "savePath", "path": "converted-layout.docx", "document": document}])
strip_snapshot("converted-layout.docx", "converted-layout-standard.docx")

standard = Document(OUT / "converted-layout-standard.docx")
section = standard.sections[0]
assert section.page_width > section.page_height
for actual, expected in [(section.top_margin.mm, 24), (section.bottom_margin.mm, 31),
                         (section.left_margin.mm, 17), (section.right_margin.mm, 27)]:
    near(actual, expected)
pa, pb = [p.paragraph_format for p in standard.paragraphs]
assert pa.space_before.pt == 12.5 and pa.space_after.pt == 6.5
assert pa.left_indent.pt == 36 and pa.first_line_indent.pt == -18
assert pa.keep_with_next is False and pa.keep_together is False and pa.page_break_before is False
assert pb.first_line_indent.pt == 24 and pb.keep_with_next is True
assert pb.keep_together is True and pb.page_break_before is True
reopened = rpc([{"command": "openPath", "path": "converted-layout-standard.docx"}])[0]
assert not reopened["protected"], reopened["warnings"]
for original, actual in zip(document["content"], reopened["document"]["content"]):
    for key in ["spaceBefore", "spaceAfter", "firstLineIndent", "keepNext", "keepLines", "pageBreakBefore"]:
        assert original["attrs"].get(key) == actual["attrs"].get(key), (key, original, actual)

# UI-created documents use numeric millimetres/points and must produce one ind tag per pPr.
ui_doc = {"type": "doc", "attrs": {"orientation": "portrait", "marginTop": 5, "marginBottom": 50,
                                     "marginLeft": 19.5, "marginRight": 22.5}, "content": [
    {"type": "heading", "attrs": {"level": 1, "spaceBefore": 0, "spaceAfter": 8,
                                    "firstLineIndent": 0, "keepNext": False, "keepLines": False,
                                    "pageBreakBefore": False},
     "content": [{"type": "text", "text": "标题默认关闭保持下段"}]},
    {"type": "paragraph", "attrs": {"indent": 3, "spaceBefore": 9, "spaceAfter": 18,
                                      "firstLineIndent": -24, "keepNext": True, "keepLines": True,
                                      "pageBreakBefore": True, "lineHeight": "2"},
     "content": [{"type": "text", "text": "悬挂缩进与分页控制"}]},
    {"type": "bulletList", "content": [{"type": "listItem", "content": [
        {"type": "paragraph", "attrs": {"firstLineIndent": 0},
         "content": [{"type": "text", "text": "默认列表继续继承编号缩进"}]}]}]},
]}
rpc([{"command": "savePath", "path": "ui-layout.docx", "document": ui_doc}])
strip_snapshot("ui-layout.docx", "ui-layout-standard.docx")
check = Document(OUT / "ui-layout-standard.docx")
assert check.paragraphs[0].paragraph_format.keep_with_next is False
assert check.paragraphs[1].paragraph_format.left_indent.pt == 54
assert check.paragraphs[1].paragraph_format.first_line_indent.pt == -24
near(check.sections[0].top_margin.mm, 5)
near(check.sections[0].bottom_margin.mm, 50)
with ZipFile(OUT / "ui-layout-standard.docx") as z:
    xml = etree.fromstring(z.read("word/document.xml"))
    for ppr in xml.iter(W + "pPr"):
        assert len(ppr.findall(W + "ind")) <= 1
    list_p = list(xml.iter(W + "p"))[-1]
    assert list_p.find(W + "pPr").find(W + "ind") is None

# Unsupported mixed sections and extreme margins must protect the original file.
mixed = Document()
configure_a4(mixed.sections[0])
mixed.add_paragraph("第一节")
second = mixed.add_section(WD_SECTION_START.NEW_PAGE)
configure_a4(second, landscape=True)
second.top_margin = Mm(2)
mixed.add_paragraph("第二节")
mixed.save(OUT / "mixed-layout.docx")
protected = rpc([{"command": "openPath", "path": "mixed-layout.docx"}])[0]
assert protected["protected"] and any("多节" in w for w in protected["warnings"])
blocked = rpc([{"command": "openPath", "path": "mixed-layout.docx"},
               {"command": "savePath", "path": "mixed-layout.docx", "document": protected["document"]}],
              allow_errors=True)
assert "error" in blocked[-1], blocked
extreme = Document()
configure_a4(extreme.sections[0])
extreme.sections[0].left_margin = Mm(60)
extreme.add_paragraph("超出支持范围的页边距")
extreme.save(OUT / "extreme-layout.docx")
limited = rpc([{"command": "openPath", "path": "extreme-layout.docx"}])[0]
assert limited["protected"] and any("页边距超出" in w for w in limited["warnings"])
assert limited["document"]["attrs"]["marginLeft"] == 50
print("独立 DOCX 布局验证通过：页边距、段距、缩进、样式继承、显式关闭分页开关、无快照重开、多节原文件保护。")
