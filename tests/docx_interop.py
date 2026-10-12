"""Independent DOCX validation using python-docx, not the Rust snapshot."""
import json
from pathlib import Path
import subprocess
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'test-output'
OUT.mkdir(exist_ok=True)
Image.new('RGB',(240,140),(35,91,171)).save(OUT/'image.png')
d=Document()
d.add_heading('外部 Word 标题',0)
p=d.add_paragraph('正常文字，')
r=p.add_run('加粗中文');r.bold=True;r.font.size=Pt(18);r.font.color.rgb=RGBColor.from_string('235BAB')
d.add_paragraph('第一项',style='List Bullet')
t=d.add_table(rows=3,cols=3);t.cell(0,0).merge(t.cell(0,1)).text='合并单元格'
t.cell(1,0).merge(t.cell(2,0)).text='跨行单元格'
t.cell(1,1).text='表格正文'
d.add_picture(str(OUT/'image.png'),width=Inches(2.5))
d.sections[0].header.paragraphs[0].text='特殊页眉'
d.save(OUT/'foreign.docx')

def rpc(reqs):
    binary=ROOT/'target/debug/rust-word'
    data='\n'.join(json.dumps(dict(id=i,**r),ensure_ascii=False) for i,r in enumerate(reqs))+'\n'
    process=subprocess.run([str(binary),'--rpc',str(OUT)],input=data,text=True,capture_output=True,check=True)
    result=[json.loads(line) for line in process.stdout.splitlines()]
    for r in result: assert 'error' not in r,r
    return [r['data'] for r in result]

def verify():
    path=OUT/'ui-roundtrip.docx'
    d=Document(path)
    assert '中文实际输入' in '\n'.join(p.text for p in d.paragraphs)
    assert len(d.tables)==1 and len(d.tables[0].rows)==3
    assert len(d.inline_shapes)==1
    assert any(r.bold for p in d.paragraphs for r in p.runs)
    assert d.sections[0].page_width>d.sections[0].page_height
    # Delete internal snapshot to prove the standard OOXML itself is sufficient.
    with ZipFile(path) as source,ZipFile(OUT/'without-snapshot.docx','w',ZIP_DEFLATED) as dest:
        for n in source.namelist():
            if n!='word/rustword.json':dest.writestr(n,source.read(n))
    reopened=rpc([{'command':'openPath','path':'without-snapshot.docx'}])[0]
    flat=json.dumps(reopened['document'],ensure_ascii=False)
    assert '中文实际输入' in flat and '替换完成' in flat and 'image/png;base64' in flat and 'table' in flat
    foreign=rpc([{'command':'openPath','path':'foreign.docx'}])[0]
    assert foreign['protected']
    assert any('页眉' in w for w in foreign['warnings'])
    assert 'rowspan": 2' in json.dumps(foreign['document'])
    converted=rpc([{'command':'openPath','path':'foreign.docx'}, {'command':'savePath','path':'foreign-converted.docx','document':foreign['document']}])[-1]
    c=Document(OUT/'foreign-converted.docx')
    assert len(c.tables)==1 and len(c.inline_shapes)==1 and '正常文字' in '\n'.join(p.text for p in c.paragraphs)
    print('独立 python-docx 读取、图片、表格、字体、横向纸张、无私有快照重开、外部 DOCX 转换：通过')

if __name__=='__main__':
    import sys
    if '--verify' in sys.argv:verify()
    else:print('外部 Word 测试文档已创建')
